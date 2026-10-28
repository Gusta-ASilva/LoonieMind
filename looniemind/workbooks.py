"""Planilhas operacionais geradas pela aplicação, com campos e totais rastreáveis."""
from datetime import date
from decimal import Decimal
import csv
import io

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from .domain import DETAIL_FIELDS
from .synthetic import IMPORT_HEADERS, SCENARIOS, financial_history, import_csv

TEAL, FOREST, MUTED = "258F83", "173C35", "EAF2EF"
CURRENCY = '"R$" #,##0.00;[Red]("R$" #,##0.00)'
DETAIL_LABELS = IMPORT_HEADERS[5:]
REPORT_HEADERS = ["Data", "Descrição", "Tipo", "Categoria", "Valor (R$)", *DETAIL_LABELS,
                  "Revisado", "Origem", "Importação", "Linha original", "ID lançamento"]


def safe_text(value):
    text = str(value or "")
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")) else text


def style_table(sheet, name, widths, money_columns=()):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.tabColor = TEAL
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=11, color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor=FOREST)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, index).column_letter].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="182D27")
            cell.alignment = Alignment(vertical="top")
        for column in money_columns:
            row[column-1].number_format = CURRENCY
    sheet.auto_filter.ref = sheet.dimensions
    if sheet.max_row > 1:
        table = Table(displayName=name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)


def serialized(book):
    book.calculation.fullCalcOnLoad = True
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def report_rows(records, excel=False):
    return [[date.fromisoformat(r["data"]) if excel else r["data"], safe_text(r["descricao"]),
             {1: "Receita", 2: "Despesa"}.get(r["tipo"], "Pendente"), safe_text(r["categoria"]),
             Decimal(r["valor_centavos"]) / 100, *(safe_text(r.get(field)) for field in DETAIL_FIELDS),
             "Sim" if r["revisado"] else "Não", safe_text(r["origem"]), r["id_planilha"] or "",
             r["linha_origem"] or "", r["id_lancamento"]] for r in records]


def report_csv(records):
    output = io.StringIO(newline="")
    writer = csv.writer(output, delimiter=";", lineterminator="\r\n")
    writer.writerow(REPORT_HEADERS)
    for row in report_rows(records):
        row[4] = f"{row[4]:.2f}".replace(".", ",")
        writer.writerow(row)
    return output.getvalue().encode("utf-8-sig")


def import_template(file_format="xlsx"):
    samples = financial_history(SCENARIOS[0], days=90)[-240:]
    if file_format == "csv":
        return import_csv(samples)
    book = Workbook()
    input_sheet = book.active
    input_sheet.title = "Movimentacoes"
    input_sheet.append(IMPORT_HEADERS)
    # A entrada começa vazia: exemplos fictícios ficam em outra aba.
    style_table(input_sheet, "Entrada", [14, 48, 18, 16, 24, 25, 28, 23, 23, 24, 16, 55, 28], [3])
    for row in range(2, 1002):
        input_sheet.cell(row, 1).number_format = "dd/mm/yyyy"
        input_sheet.cell(row, 3).number_format = CURRENCY
    for column, choices in [("D", '"Receita,Despesa"'), ("J", '"PIX,Transferência,Cartão,Boleto,Dinheiro,Débito"')]:
        validation = DataValidation(type="list", formula1=choices, allow_blank=True)
        validation.errorTitle, validation.error = "Opção inválida", "Escolha um item da lista."
        validation.showErrorMessage = True
        input_sheet.add_data_validation(validation)
        validation.add(f"{column}2:{column}1001")
    amount = DataValidation(type="decimal", operator="greaterThan", formula1=0, allow_blank=True)
    amount.showErrorMessage = True
    amount.error = "Use um valor positivo. Informe a direção na coluna Tipo."
    input_sheet.add_data_validation(amount)
    amount.add("C2:C1001")
    example = book.create_sheet("Exemplo_sintetico")
    example.append(IMPORT_HEADERS)
    for r in samples:
        example.append([date.fromisoformat(r["data"]), r["descricao"], r["valor_centavos"]/100,
            "Receita" if r["tipo"] == 1 else "Despesa", r["categoria"], *(r[f] for f in DETAIL_FIELDS)])
    style_table(example, "Exemplos", [14, 48, 18, 16, 24, 25, 28, 23, 23, 24, 16, 55, 28], [3])
    for cell in example["A"][1:]:
        cell.number_format = "dd/mm/yyyy"
    guide = book.create_sheet("Orientacoes")
    guide.append(["Campo", "Preenchimento", "Regra e uso"])
    instructions = [
        ("Como usar", "Preencha Movimentacoes", "Aba de entrada vazia. No sistema, selecione a aba correta, confira a prévia e confirme."),
        ("Exemplo_sintetico", "240 registros fictícios", "Somente demonstração. Não misture os exemplos com movimentações reais."),
        ("Data", "Obrigatório", "Data em que o dinheiro entrou ou saiu; dd/mm/aaaa. Não aceita data futura."),
        ("Descrição", "Obrigatório", "Texto da operação. Também é a entrada usada pelo classificador."),
        ("Valor", "Obrigatório", "Número positivo, com até 2 casas decimais. O tipo define entrada ou saída."),
        ("Tipo", "Recomendado", "Receita ou Despesa. Sem tipo, o valor fica fora do saldo até a revisão."),
        ("Categoria", "Recomendado", "Com tipo e categoria válidos, a importação fica revisada. Vazio exige revisão."),
        ("Documento", "Opcional · 80 caracteres", "Nota, recibo ou referência documental."),
        ("Contraparte", "Opcional · 120 caracteres", "Cliente, fornecedor ou instituição envolvida."),
        ("Centro de custo", "Opcional · 80 caracteres", "Área responsável, por exemplo Loja ou Administrativo."),
        ("Conta", "Opcional · 80 caracteres", "Conta ou caixa de origem. O saldo do dashboard é consolidado da empresa."),
        ("Forma de pagamento", "Opcional · 60 caracteres", "PIX, transferência, cartão, boleto, dinheiro ou outra descrição."),
        ("Competência", "Opcional", "Mês econômico, aaaa-mm ou mm/aaaa. O caixa usa Data, não Competência."),
        ("Observações", "Opcional · 500 caracteres", "Contexto complementar da movimentação."),
        ("ID externo", "Opcional · 100 caracteres", "Código do sistema de origem, preservado para rastreamento; não é uma chave de deduplicação."),
        ("Limites", "10 MB e 10.000 linhas", "CSV, XLS ou XLSX. Divida arquivos maiores por ano, conta ou período."),
        ("Duplicidades", "Confira os avisos", "Reimportação do mesmo arquivo/aba é bloqueada; linhas semelhantes exigem confirmação."),
    ]
    for row in instructions:
        guide.append(row)
    style_table(guide, "Guia", [25, 34, 115])
    for row in guide.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
        guide.row_dimensions[row[0].row].height = 30
    return serialized(book)


def finance_report(records, company, filters, current_balance):
    book = Workbook()
    summary = book.active
    summary.title = "Resumo"
    detail = book.create_sheet("Lancamentos")
    detail.append(REPORT_HEADERS)
    for row in report_rows(records, excel=True):
        detail.append(row)
    style_table(detail, "Movimentos", [14, 48, 16, 24, 20, 26, 28, 24, 24, 24, 16, 55, 28, 14, 25, 38, 18, 18], [5])
    for cell in detail["A"][1:]:
        cell.number_format = "dd/mm/yyyy"
    end = max(2, len(records)+1)
    kinds, amounts = f"Lancamentos!C2:C{end}", f"Lancamentos!E2:E{end}"
    summary.append(["LOONIEMIND · RESUMO FINANCEIRO", "Valor", "Contexto"])
    summary.append(["Empresa", safe_text(company["nome"]), "Dados fictícios" if company["demo"] else "Dados da empresa"])
    summary.append(["Gerado em", date.today(), "Fotografia dos dados na data da exportação"])
    summary.append(["Filtros", json_filters(filters), "Todos os totais abaixo usam somente os lançamentos exportados, salvo saldo consolidado"])
    summary.append(["Lançamentos exportados", len(records), "Inclui registros pendentes"])
    summary.append(["Receitas do recorte", f'=SUMIF({kinds},"Receita",{amounts})', "Valores com tipo Receita"])
    summary.append(["Despesas do recorte", f'=SUMIF({kinds},"Despesa",{amounts})', "Valores com tipo Despesa"])
    summary.append(["Resultado do recorte", "=B6-B7", "Receitas menos despesas; não é o saldo disponível"])
    summary.append(["Valor sem tipo", f'=SUMIF({kinds},"Pendente",{amounts})', "Fora dos totais e do saldo até a definição do tipo"])
    summary.append(["Revisões pendentes", sum(not r["revisado"] for r in records), "Revise na tela Classificação"])
    summary.append(["Saldo consolidado atual", current_balance/100, "Saldo inicial + todas as receitas - despesas desde a referência; independe dos filtros deste relatório"])
    summary.append(["Saldo inicial da empresa", company["saldo_inicial_centavos"]/100, "Referência em " + company["data_saldo"]])
    summary.append(["Critério", "Regime de caixa", "Data representa realização. Competência é apenas um atributo informativo."])
    style_table(summary, "ResumoFinanceiro", [43, 42, 110], [2])
    summary["B3"].number_format = "dd/mm/yyyy"
    summary["B5"].number_format = summary["B10"].number_format = "0"
    for row in (6, 7, 8, 9, 11, 12):
        summary.cell(row, 2).number_format = CURRENCY
    for row in summary.iter_rows(min_row=2):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
        summary.row_dimensions[row[0].row].height = 31
    month = book.create_sheet("Mensal")
    month.append(["Mês", "Receitas", "Despesas", "Resultado", "Movimentações"])
    months = sorted({r["data"][:7] for r in records})
    for index, prefix in enumerate(months, 2):
        begin = date.fromisoformat(prefix+"-01")
        until = date(begin.year + (begin.month == 12), begin.month % 12 + 1, 1)
        criterion = f'">="&DATE({begin.year},{begin.month},1)'
        upper = f'"<"&DATE({until.year},{until.month},1)'
        dates = f"Lancamentos!A2:A{end}"
        month.append([prefix, f'=SUMIFS({amounts},{kinds},"Receita",{dates},{criterion},{dates},{upper})',
            f'=SUMIFS({amounts},{kinds},"Despesa",{dates},{criterion},{dates},{upper})',
            f"=B{index}-C{index}", f'=COUNTIFS({dates},{criterion},{dates},{upper})'])
    style_table(month, "MensalFinanceiro", [18, 22, 22, 22, 22], [2, 3, 4])
    for cell in month["A"][1:]:
        cell.number_format = "mmm/yyyy"
    if months:
        chart = BarChart()
        chart.title, chart.y_axis.title = "Entradas e saídas por mês", "R$"
        chart.add_data(Reference(month, min_col=2, max_col=3, min_row=1, max_row=len(months)+1), titles_from_data=True)
        chart.set_categories(Reference(month, min_col=1, min_row=2, max_row=len(months)+1))
        chart.width, chart.height = 24, 12
        month.add_chart(chart, "G2")
    for sheet_name, key, column in [("Categorias", "categoria", "D"), ("Centros_de_custo", "centro_custo", "H")]:
        sheet = book.create_sheet(sheet_name)
        sheet.append([sheet_name.replace("_", " "), "Receitas", "Despesas", "Resultado", "Movimentações"])
        groups = sorted({r.get(key) or "" for r in records})
        source = f"Lancamentos!{column}2:{column}{end}"
        for index, value in enumerate(groups, 2):
            # Prefix '=' forces an exact text criterion; wildcard characters are escaped.
            literal = safe_text(value).replace("~", "~~").replace("*", "~*").replace("?", "~?").replace('"', '""')
            crit = f'"={literal}"'
            sheet.append([safe_text(value) or "(Não informado)", f'=SUMIFS({amounts},{kinds},"Receita",{source},{crit})',
                f'=SUMIFS({amounts},{kinds},"Despesa",{source},{crit})', f"=B{index}-C{index}", f'=COUNTIF({source},{crit})'])
        style_table(sheet, sheet_name.replace("_", ""), [32, 22, 22, 22, 22], [2, 3, 4])
    book._sheets = [summary, month, book["Categorias"], book["Centros_de_custo"], detail]
    return serialized(book)


def json_filters(filters):
    labels = {"inicio": "De", "fim": "Até", "tipo": "Tipo", "categoria": "Categoria", "busca": "Busca", "pendentes": "Pendentes"}
    return " · ".join(f"{labels.get(k,k)}: {v}" for k, v in (filters or {}).items() if v not in (None, "", False)) or "Todo o histórico"


def forecast_report(result, company):
    from openpyxl.chart import LineChart
    book = Workbook()
    summary = book.active
    summary.title = "Resumo"
    summary.append(["LOONIEMIND · PROJEÇÃO", "Valor", "Contexto"])
    for row in [
        ["Empresa", safe_text(company["nome"]), "Dados sintéticos" if result["sintetico"] else "Histórico da empresa"],
        ["Base até", result["origem_previsao"], f"{result['dias_historico']} dias de histórico"],
        ["Horizonte", result["dias"], "Dias a partir do dia seguinte à data-base"],
        ["Saldo atual", result["saldo_atual_centavos"]/100, "Saldo realizado de partida"],
        ["Receitas estimadas", f"=SUM(Previsao!B2:B{len(result['serie'])+1})", "Estimativa estatística"],
        ["Despesas estimadas", f"=SUM(Previsao!C2:C{len(result['serie'])+1})", "Estimativa estatística"],
        ["Saldo final estimado", f"=Previsao!E{len(result['serie'])+1}", "Saldo atual + entradas - saídas estimadas"],
        ["Menor saldo estimado", f"=MIN(Previsao!E2:E{len(result['serie'])+1})", "Menor saldo entre as datas futuras"],
        ["Método", result["metodo"], "Seleção por menor MAE na validação cronológica"],
        ["Limitação", result["aviso"], "Não representa agenda de contas a pagar nem garantia de caixa futuro"],
    ]:
        summary.append(row)
    style_table(summary, "ResumoPrevisao", [38, 50, 100])
    for index in range(5, 10):
        summary.cell(index, 2).number_format = CURRENCY
    for row in summary.iter_rows(min_row=2):
        row[1].alignment = row[2].alignment = Alignment(wrap_text=True, vertical="top")
        summary.row_dimensions[row[0].row].height = 36 if row[0].row != 11 else 65
    daily = book.create_sheet("Previsao")
    daily.append(["Data", "Receitas estimadas", "Despesas estimadas", "Resultado estimado", "Saldo estimado", "Data no gráfico"])
    for index, row in enumerate(result["serie"], 2):
        opening = "Resumo!B5" if index == 2 else f"E{index-1}"
        daily.append([date.fromisoformat(row["data"]), row["receitas_centavos"]/100, row["despesas_centavos"]/100,
                      f"=B{index}-C{index}", f"={opening}+D{index}", date.fromisoformat(row["data"]).strftime("%d/%m")])
    style_table(daily, "ProjecaoDiaria", [16, 23, 23, 23, 23, 18], [2, 3, 4, 5])
    for cell in daily["A"][1:]:
        cell.number_format = "dd/mm/yyyy"
    chart = LineChart()
    chart.title, chart.y_axis.title = "Evolução do saldo estimado", "R$"
    chart.add_data(Reference(daily, min_col=5, min_row=1, max_row=len(result["serie"])+1), titles_from_data=True)
    chart.set_categories(Reference(daily, min_col=6, min_row=2, max_row=len(result["serie"])+1))
    chart.width, chart.height = 24, 12
    daily.add_chart(chart, "G2")
    methods = book.create_sheet("Metodos")
    methods.append(["Série", "Método escolhido", "MAE LightGBM (R$/dia)", "MAE referência (R$/dia)", "WAPE escolhido", "Dias de validação", "Início", "Fim"])
    for name, model in result["modelos"].items():
        evaluation = model.get("avaliacao") or {}
        methods.append([name.title(), model["metodo"], (evaluation.get("mae_modelo_centavos", 0)/100) if evaluation else None,
            (evaluation.get("mae_referencia_centavos", 0)/100) if evaluation else None,
            evaluation.get("wape"), evaluation.get("dias_teste"), evaluation.get("inicio_teste"), evaluation.get("fim_teste")])
    style_table(methods, "MetodosPrevisao", [20, 40, 27, 27, 22, 24, 17, 17], [3, 4])
    for cell in methods["E"][1:]:
        cell.number_format = "0.0%"
    return serialized(book)
