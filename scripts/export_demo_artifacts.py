"""Extrai os relatórios reais da aplicação e dados para a documentação desta versão."""
import json
from pathlib import Path
from datetime import date, datetime

from openpyxl import load_workbook

from looniemind.db import Database
from looniemind.models import ModelService, ForecastService
from looniemind.services import FinanceService
from looniemind.workbooks import import_template, forecast_report


def main():
    directory = Path('.build/v2')
    directory.mkdir(parents=True, exist_ok=True)
    db = Database()
    db.initialize()
    finance = FinanceService(db, ModelService(db), ForecastService(db))
    with db.connect() as conn:
        identifier = conn.execute("SELECT id_empresa FROM Empresas WHERE descricao='Base sintética v2 · aurora'").fetchone()[0]
    forecast = finance.forecasts.calculate(identifier)
    artifacts = {
        'Modelo_Importacao': import_template(),
        'Relatorio_Demonstracao': finance.export(identifier, {}, 'xlsx'),
        'Projecao_Demonstracao': forecast_report(forecast, finance.get_company(identifier)),
    }
    specifications = []
    for name, content in artifacts.items():
        path = directory / (name + '-runtime.xlsx')
        path.write_bytes(content)
        book = load_workbook(path)
        sheets = []
        for sheet in book:
            rows = []
            # Formatted blank input rows do not become source records.
            max_row = 1 if sheet.title == 'Movimentacoes' else sheet.max_row
            for row in sheet.iter_rows(max_row=max_row):
                rows.append([{'date': cell.value.date().isoformat() if isinstance(cell.value, datetime) else cell.value.isoformat()}
                             if isinstance(cell.value, (date, datetime)) else cell.value for cell in row])
            sheets.append({'name': sheet.title, 'rows': rows,
                'widths': {key: value.width for key, value in sheet.column_dimensions.items()},
                'row_heights': {key: value.height for key, value in sheet.row_dimensions.items() if value.height},
                'money_columns': [cell.column for cell in sheet[2] if 'R$' in cell.number_format] if sheet.max_row > 1 else [],
                'date_columns': [cell.column for cell in sheet[2] if cell.is_date] if sheet.max_row > 1 else []})
        specifications.append({'name': name, 'sheets': sheets})
    (directory/'workbooks.json').write_text(json.dumps(specifications, ensure_ascii=False, default=str), encoding='utf-8')
    print(json.dumps({a['name']: [(s['name'], len(s['rows'])) for s in a['sheets']] for a in specifications}, ensure_ascii=False))


if __name__ == '__main__':
    main()
