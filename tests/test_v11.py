from datetime import date, timedelta
import io
import sqlite3

from openpyxl import load_workbook

from looniemind.db import Database
from looniemind.domain import DETAIL_FIELDS
from looniemind.models import TRAINING_TOPICS
from looniemind.synthetic import classification_data


def test_rich_import_export_roundtrip_and_partial_update(client, company):
    values = [date.today().isoformat(), "Compra documentada", "123,45", "Despesa", "Fornecedores", "NF-42", "Fornecedor Modelo",
              "Operação", "Banco Modelo", "PIX", "10/2026", "Observação preservada", "EXT-1"]
    headers = "Data;Descrição;Valor;Tipo;Categoria;Documento;Contraparte;Centro de custo;Conta;Forma de pagamento;Competência;Observações;ID externo"
    body = (headers+"\n"+";".join(values)).encode("utf-8-sig")
    prefix = f"/api/companies/{company}"
    uploaded = client.post(prefix+"/imports", files={"arquivo": ("dados.csv", body)}).json()
    assert uploaded["invalidos"] == 0 and uploaded["validos"] == 1
    assert client.post(prefix+f"/imports/{uploaded['id_planilha']}/confirm", json={}).status_code == 200
    row = client.get(prefix+"/transactions").json()["itens"][0]
    assert row["documento"] == "NF-42" and row["competencia"] == "2026-10" and row["contraparte"] == "Fornecedor Modelo"
    assert client.get(prefix+"/transactions?busca=NF-42").json()["total"] == 1
    payload = {k: row[k] for k in ("descricao", "data", "tipo", "categoria", "valor_centavos")}
    assert client.put(prefix+f"/transactions/{row['id_lancamento']}", json=payload).status_code == 200
    assert client.get(prefix+"/transactions").json()["itens"][0]["id_externo"] == "EXT-1"
    exported = client.get(prefix+"/export?formato=csv").content
    other = client.post("/api/companies", json={"nome": "Destino"}).json()["id_empresa"]
    preview = client.post(f"/api/companies/{other}/imports", files={"arquivo": ("reimport.csv", exported)}).json()
    assert preview["validos"] == 1 and preview["linhas"][0]["centro_custo"] == "Operação"


def test_template_separates_blank_input_from_examples(client, company):
    content = client.get("/api/template?formato=xlsx").content
    book = load_workbook(io.BytesIO(content))
    assert book.sheetnames == ["Movimentacoes", "Exemplo_sintetico", "Orientacoes"]
    assert book["Movimentacoes"]["B2"].value is None
    assert book["Movimentacoes"].max_column == 13
    assert book["Exemplo_sintetico"].max_row == 241
    response = client.post(f"/api/companies/{company}/imports", files={"arquivo": ("modelo.xlsx", content)}).json()
    assert response["validos"] == 0
    preview = client.post(f"/api/companies/{company}/imports/{response['id_planilha']}/preview",
                          json={"sheet": "Exemplo_sintetico"}).json()
    assert preview["validos"] == 240 and preview["invalidos"] == 0


def test_daily_balance_reconciles_with_summary(client, company, movement):
    prefix = f"/api/companies/{company}"
    for days, kind, value in [(31, 1, 90000), (2, 2, 5000), (0, 1, 10000)]:
        client.post(prefix+"/transactions", json={**movement, "data": (date.today()-timedelta(days=days)).isoformat(),
            "tipo": kind, "valor_centavos": value})
    data = client.get(prefix+"/summary?periodo=30d").json()
    assert len(data["fluxo_diario"]) == 30
    assert data["saldo_abertura_grafico_centavos"] == 190000
    assert data["fluxo_diario"][-1]["saldo_centavos"] == data["saldo_centavos"] == 195000
    assert data["resultado_centavos"] == 5000


def test_schema_migration_keeps_v1_data_and_snapshot(tmp_path):
    database = Database(tmp_path)
    database.initialize()
    with sqlite3.connect(database.path) as connection:
        for field in DETAIL_FIELDS:
            connection.execute(f"ALTER TABLE Lancamentos DROP COLUMN {field}")
        connection.execute("PRAGMA user_version=1")
        connection.execute("DELETE FROM schema_migrations WHERE version=2")
    database.initialize()
    with database.connect() as connection:
        assert connection.execute("PRAGMA user_version").fetchone()[0] == 2
        assert set(DETAIL_FIELDS) <= {r[1] for r in connection.execute("PRAGMA table_info(Lancamentos)")}
    assert (tmp_path / "backups/antes-da-versao-2.sqlite3").is_file()


def test_large_dataset_has_balanced_labels_and_family_markers():
    rows = classification_data(TRAINING_TOPICS)
    assert len(rows) == 14400
    assert len({r["grupo"] for r in rows}) == 288
    assert len({r["id_exemplo"] for r in rows}) == len(rows)
