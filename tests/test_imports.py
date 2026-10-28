from datetime import date, timedelta
import io

from openpyxl import Workbook


def upload(client, company, content, name="movimentos.csv"):
    return client.post(f"/api/companies/{company}/imports", files={"arquivo": (name, content)})


def data_csv(rows):
    return ("Data;Descrição;Valor;Tipo;Categoria\n" + "\n".join(rows)).encode("utf-8-sig")


def test_preview_confirm_idempotence_and_undo(client, company):
    today = date.today().strftime("%d/%m/%Y")
    content = data_csv([f"{today};Compra fornecedor;100,20;Despesa;Fornecedores", f"{today};Venda de produtos;150,10;Receita;Vendas"])
    preview = upload(client, company, content).json()
    assert preview["validos"] == 2 and preview["invalidos"] == 0
    assert client.get(f"/api/companies/{company}/transactions").json()["total"] == 0
    path = f"/api/companies/{company}/imports/{preview['id_planilha']}"
    assert client.post(path + "/confirm", json={}).json()["importados"] == 2
    assert client.post(path + "/confirm", json={}).json()["ja_confirmado"] is True
    assert client.get(f"/api/companies/{company}/summary").json()["saldo_centavos"] == 104990
    duplicate = upload(client, company, content).json()
    assert duplicate["duplicado"] is True
    assert client.post(f"/api/companies/{company}/imports/{duplicate['id_planilha']}/confirm", json={}).status_code == 409
    assert client.get(f"/api/companies/{company}/transactions").json()["total"] == 2
    assert client.post(path + "/undo").json()["removidos"] == 2
    assert client.get(f"/api/companies/{company}/transactions").json()["total"] == 0
    assert client.post(path + "/confirm", json={}).status_code == 409


def test_partial_import_requires_consent_and_unknown_type_excluded(client, company):
    today = date.today().isoformat()
    content = data_csv([f"{today};Sem tipo;50,00;;", f"{today};Inválido;texto;Receita;Vendas"])
    preview = upload(client, company, content).json()
    assert preview["validos"] == 1 and preview["invalidos"] == 1 and preview["sem_tipo"] == 1
    path = f"/api/companies/{company}/imports/{preview['id_planilha']}/confirm"
    assert client.post(path, json={}).status_code == 422
    assert client.post(path, json={"permitir_parcial": True}).json()["importados"] == 1
    summary = client.get(f"/api/companies/{company}/summary").json()
    assert summary["saldo_centavos"] == 100000 and summary["sem_tipo"] == 1
    exported = client.get(f"/api/companies/{company}/export?formato=csv").content
    other = client.post("/api/companies", json={"nome": "Reimportação pendente"}).json()["id_empresa"]
    pending_preview = upload(client, other, exported).json()
    assert pending_preview["validos"] == 1 and pending_preview["sem_tipo"] == 1


def test_ambiguous_money_requires_explicit_format(client, company):
    preview = upload(client, company, data_csv([f"{date.today()};Produto;1.234;Receita;Vendas"])).json()
    assert preview["invalidos"] == 1
    updated = client.post(f"/api/companies/{company}/imports/{preview['id_planilha']}/preview", json={"decimal_format": "br"}).json()
    assert updated["linhas"][0]["valor_centavos"] == 123400


def test_excel_sheet_selection_and_debit_credit(client, company):
    book = Workbook()
    book.active.title = "Capa"
    book.active.append(["Descrição da planilha"])
    sheet = book.create_sheet("Extrato")
    sheet.append(["Data", "Histórico", "Crédito", "Débito"])
    sheet.append([date.today().isoformat(), "Cliente", "200,00", ""])
    sheet.append([date.today().isoformat(), "Material", "", "50,00"])
    buffer = io.BytesIO()
    book.save(buffer)
    preview = upload(client, company, buffer.getvalue(), "arquivo.xlsx").json()
    assert preview["campos_ausentes"]
    updated = client.post(f"/api/companies/{company}/imports/{preview['id_planilha']}/preview", json={"sheet": "Extrato", "mapping": None}).json()
    assert updated["validos"] == 2
    assert [row["tipo"] for row in updated["linhas"]] == [1, 2]


def test_similar_transactions_are_preserved_with_acknowledgement(client, company):
    today = date.today().isoformat()
    content = data_csv([f"{today};Venda;10,00;Receita;Vendas", f"{today};Venda;10,00;Receita;Vendas"])
    p = upload(client, company, content).json()
    assert client.post(f"/api/companies/{company}/imports/{p['id_planilha']}/confirm", json={}).status_code == 200
    other = upload(client, company, content + b"\n").json()
    assert other["possiveis_duplicados"] == 2
    path = f"/api/companies/{company}/imports/{other['id_planilha']}/confirm"
    assert client.post(path, json={}).status_code == 409
    assert client.post(path, json={"manter_semelhantes": True}).status_code == 200
    assert client.get(f"/api/companies/{company}/transactions").json()["total"] == 4


def test_failed_confirmation_rolls_back_all_rows(client, company):
    today = date.today().isoformat()
    p = upload(client, company, data_csv([f"{today};OK;10,00;Receita;Vendas", f"{today};FAIL;10,00;Receita;Vendas"])).json()
    with client.app.state.db.connect(write=True) as conn:
        conn.execute("CREATE TRIGGER reject_test BEFORE INSERT ON Lancamentos WHEN NEW.descricao='FAIL' BEGIN SELECT RAISE(ABORT,'test'); END")
    assert client.post(f"/api/companies/{company}/imports/{p['id_planilha']}/confirm", json={}).status_code == 409
    assert client.get(f"/api/companies/{company}/transactions").json()["total"] == 0
    assert client.get(f"/api/companies/{company}/imports").json()[0]["status"] == "preview"


def test_invalid_and_future_files_and_company_isolation(client, company):
    assert upload(client, company, b"malicious", "script.exe").status_code == 400
    assert upload(client, company, b"not an xlsx", "bad.xlsx").status_code == 400
    tomorrow = (date.today() + timedelta(days=1)).isoformat()
    p = upload(client, company, data_csv([f"{tomorrow};Futuro;10,00;Receita;Vendas"])).json()
    assert p["invalidos"] == 1
    second = client.post("/api/companies", json={"nome": "Outra empresa"}).json()["id_empresa"]
    assert client.post(f"/api/companies/{second}/imports/{p['id_planilha']}/confirm", json={}).status_code == 404


def test_model_templates_can_be_imported(client, company):
    for format in ["csv", "xlsx"]:
        template = client.get(f"/api/template?formato={format}")
        preview = upload(client, company, template.content, f"template.{format}").json()
        assert preview["validos"] == (240 if format == "csv" else 0) and preview["invalidos"] == 0


def test_legacy_excel_xls_and_cp1252(client, company):
    import xlwt
    book = xlwt.Workbook()
    sheet = book.add_sheet("Extrato")
    for row_index, row in enumerate([["Data", "Descrição", "Valor", "Tipo", "Categoria"],
        [date.today().strftime("%d/%m/%Y"), "Serviço de manutenção", "42,50", "Receita", "Serviços"]]):
        for col_index, value in enumerate(row):
            sheet.write(row_index, col_index, value)
    output = io.BytesIO()
    book.save(output)
    result = upload(client, company, output.getvalue(), "antigo.xls")
    assert result.status_code == 201 and result.json()["validos"] == 1
    content = f"Data;Descrição;Valor;Tipo\n{date.today()};Café da manhã;25,30;Despesa".encode("cp1252")
    preview = upload(client, company, content).json()
    assert preview["metadados"]["encoding"] == "cp1252"
    assert preview["linhas"][0]["descricao"] == "Café da manhã"
