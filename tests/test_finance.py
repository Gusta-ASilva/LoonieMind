from datetime import date, timedelta
import io
import json
import sqlite3
import zipfile

from openpyxl import load_workbook
import pytest

from looniemind.domain import money_to_cents


@pytest.mark.parametrize("value,style,result", [("1.234,56", "auto", 123456), ("1234.56", "auto", 123456),
    ("1,234.56", "us", 123456), ("(22,90)", "br", -2290), ("0,10", "auto", 10), ("1.234", "br", 123400)])
def test_money(value, style, result):
    assert money_to_cents(value, style) == result


@pytest.mark.parametrize("value", ["1.234", "NaN", "Infinity", "=1+1", "zero", "0", "1,2,3"])
def test_invalid_money(value):
    with pytest.raises(ValueError):
        money_to_cents(value)


def test_installation_and_local_boundary(client):
    assert client.get("/api/health").json()["status"] == "ok"
    assert client.get("/").status_code == 200
    assert client.get("/assets/app.js").status_code == 200
    assert client.get("/api/health", headers={"host": "malicious.example"}).status_code == 400
    response = client.post("/api/companies", json={"nome": "Malicioso"}, headers={"origin": "https://malicious.example"})
    assert response.status_code == 403
    assert client.get("/api/companies").json() == []


def test_precision_isolation_and_manual_audit(client, company, movement):
    second = client.post("/api/companies", json={"nome": "Outra empresa"}).json()["id_empresa"]
    created = client.post(f"/api/companies/{company}/transactions", json=movement).json()
    identifier = created["id_lancamento"]
    assert client.get(f"/api/companies/{second}/transactions").json()["total"] == 0
    assert client.put(f"/api/companies/{second}/transactions/{identifier}", json=movement).status_code == 404
    assert client.delete(f"/api/companies/{second}/transactions/{identifier}").status_code == 404
    for _ in range(10):
        client.post(f"/api/companies/{company}/transactions", json={**movement, "tipo": 1, "valor_centavos": 10, "categoria": "Vendas"})
    summary = client.get(f"/api/companies/{company}/summary").json()
    assert summary["saldo_centavos"] == 100000 - 12345 + 100
    assert client.put(f"/api/companies/{company}/transactions/{identifier}", json={**movement, "valor_centavos": 10000}).status_code == 200
    assert client.get(f"/api/companies/{company}/summary").json()["saldo_centavos"] == 90100
    audit = client.get(f"/api/companies/{company}/audit").json()
    assert audit[0]["acao"] == "corrigir_lancamento"
    assert audit[0]["detalhes"]["antes"]["valor_centavos"] == 12345


def test_initial_balance_cutoff_and_future_rejection(client, company, movement):
    yesterday = (date.today() - timedelta(days=1)).isoformat()
    client.post(f"/api/companies/{company}/transactions", json={**movement, "data": yesterday})
    client.put(f"/api/companies/{company}", json={"nome": "Empresa de teste", "saldo_inicial_centavos": 50000, "data_saldo": date.today().isoformat()})
    assert client.get(f"/api/companies/{company}/summary").json()["saldo_centavos"] == 50000
    tomorrow = (date.today() + timedelta(days=1)).isoformat()
    assert client.post(f"/api/companies/{company}/transactions", json={**movement, "data": tomorrow}).status_code == 422
    assert client.post(f"/api/companies/{company}/transactions", json={**movement, "valor_centavos": 1.5}).status_code == 422
    assert client.post(f"/api/companies/{company}/transactions", json={**movement, "valor_centavos": -1}).status_code == 422


def test_exports_filters_and_formula_protection(client, company, movement):
    client.post(f"/api/companies/{company}/transactions", json={**movement, "descricao": "=HYPERLINK(\"https://example.com\")"})
    client.post(f"/api/companies/{company}/transactions", json={**movement, "tipo": 1, "categoria": "Vendas"})
    prefix = f"/api/companies/{company}/export"
    csv = client.get(prefix + "?formato=csv&tipo=2")
    assert csv.status_code == 200 and csv.content.startswith(b"\xef\xbb\xbf")
    assert "'=HYPERLINK" in csv.content.decode("utf-8-sig")
    assert "123,45" in csv.content.decode("utf-8-sig")
    xlsx = client.get(prefix + "?formato=xlsx&tipo=2")
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook.sheetnames == ["Resumo", "Mensal", "Categorias", "Centros_de_custo", "Lancamentos"]
    sheet = workbook["Lancamentos"]
    assert sheet.max_row == 2
    assert sheet["B2"].data_type == "s" and sheet["B2"].value.startswith("'=")
    assert sheet["E2"].value == 123.45
    assert client.get(f"/api/companies/{company}/transactions?busca=%25").json()["total"] == 0
    assert client.get(f"/api/companies/{company}/transactions?inicio=2026-12-01&fim=2026-01-01").status_code == 400


def test_backup_contains_consistent_database(client, company, movement):
    client.post(f"/api/companies/{company}/transactions", json=movement)
    backup = client.get("/api/backup")
    assert backup.status_code == 200
    with zipfile.ZipFile(io.BytesIO(backup.content)) as archive:
        assert json.loads(archive.read("manifest.json"))["versao"] == 2
        connection = sqlite3.connect(":memory:")
        connection.deserialize(archive.read("looniemind.sqlite3"))
        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert connection.execute("SELECT count(*) FROM Lancamentos").fetchone()[0] == 1
        connection.close()


def test_negative_current_balance_generates_an_actual_alert(client, company):
    client.put(f"/api/companies/{company}", json={"nome": "Empresa de teste", "saldo_inicial_centavos": -10000,
        "data_saldo": date.today().isoformat()})
    summary = client.get(f"/api/companies/{company}/summary").json()
    assert summary["saldo_centavos"] == -10000
    assert summary["alertas"][0]["nivel"] == "danger"
    assert summary["alertas"][0]["titulo"] == "Saldo atual negativo"
